"""Gera o relatório executivo do projeto em Word (.docx).

Documento curto, para instruir a proposta de contratação da licença do Kiro.
Sem dados sensíveis: apenas nomes de tela, integrações e números do que foi
entregue.

Uso:
    cd backend/
    python scripts/gerar_relatorio_proposta.py [--saida CAMINHO]
"""
from __future__ import annotations

import argparse
import os
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# Identidade do projeto
APP_NOME = "Processamento Sancionatório CPSAR"
AREA = "Coordenadoria de Processamento Sancionatório (CPSAR) — DETRAN-SP"
AUTOR = "Pedro Henrique Marques Silva"

AZUL = RGBColor(0x00, 0x44, 0x7F)      # primary do app
CINZA = RGBColor(0x44, 0x4A, 0x52)

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def _titulo(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = AZUL


def _bullet(doc: Document, negrito: str, resto: str = "") -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(negrito)
    r.bold = True
    if resto:
        p.add_run(resto)


def _tabela_identificacao(doc: Document, hoje: date) -> None:
    dados = [
        ("Aplicativo", APP_NOME),
        ("Área atendida", AREA),
        ("Autor", AUTOR),
        ("Data", f"{hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}"),
    ]
    tabela = doc.add_table(rows=0, cols=2)
    tabela.style = "Light List Accent 1"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for rotulo, valor in dados:
        linha = tabela.add_row().cells
        p0 = linha[0].paragraphs[0]
        r0 = p0.add_run(rotulo)
        r0.bold = True
        linha[1].paragraphs[0].add_run(valor)
    tabela.columns[0].width = Pt(120)


def gerar(caminho: str) -> str:
    hoje = date.today()
    doc = Document()

    # Margens um pouco menores, para caber em duas páginas
    for secao in doc.sections:
        secao.top_margin = Pt(50)
        secao.bottom_margin = Pt(50)
        secao.left_margin = Pt(60)
        secao.right_margin = Pt(60)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)

    # ── Cabeçalho ────────────────────────────────────────────────────────────
    cab = doc.add_paragraph()
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cab.add_run("Relatório de Desenvolvimento")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = AZUL

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"{APP_NOME}\nProposta de contratação da licença Kiro")
    r.font.size = Pt(11)
    r.font.color.rgb = CINZA

    doc.add_paragraph()
    _tabela_identificacao(doc, hoje)

    # ── Objetivo ─────────────────────────────────────────────────────────────
    _titulo(doc, "1. Objetivo")
    doc.add_paragraph(
        "Substituir o aplicativo de Processamento Sancionatório, hoje em Power Apps com dados "
        "em listas do SharePoint e integrações por fluxos do Power Automate, por uma aplicação "
        "própria — com banco de dados estruturado e integração direta com a API do SEI. "
        "O desenvolvimento é conduzido com o assistente Kiro."
    )

    # ── O que o sistema faz ──────────────────────────────────────────────────
    _titulo(doc, "2. O que o sistema faz")
    doc.add_paragraph(
        "Gerencia o ciclo de vida dos processos administrativos sancionatórios contra agentes "
        "regulados (autoescolas, empresas de vistoria, desmontes, peritos, entre outros), "
        "da chegada do relatório de fiscalização até o encerramento do processo."
    )
    _bullet(doc, "Caixa de Entrada", " — recebe automaticamente do SEI os relatórios remetidos pela fiscalização.")
    _bullet(doc, "Análise do relatório", " — leitura dos documentos, histórico e apontamentos de irregularidade apurados do checklist da fiscalização.")
    _bullet(doc, "Triagem", " — arquivamento, termo de ajustamento de conduta ou instauração de processo.")
    _bullet(doc, "Instauração automatizada", " — cria o processo no SEI, monta o conjunto probatório em PDF, gera o termo e envia ao bloco de assinatura.")
    _bullet(doc, "Processos em Andamento", " — acompanhamento das fases, prazos e documentos de cada processo.")
    _bullet(doc, "Consulta Unificada", " — visão única de relatórios e processos para consulta.")
    _bullet(doc, "Notificações e prazos", " — avisa quando o interessado se manifesta e emite certidão de decurso quando o prazo vence.")

    # ── Fases automatizadas ──────────────────────────────────────────────────
    _titulo(doc, "3. Fases do processo automatizadas")
    doc.add_paragraph(
        "Todas as fases previstas pela área foram implementadas, com avanço automático a partir "
        "dos andamentos do SEI: instauração, defesa prévia, instrução processual, alegações finais, "
        "julgamento, recurso e encerramento. O sistema detecta a manifestação do interessado, "
        "controla os prazos legais e gera as certidões correspondentes."
    )

    # ── Integrações ──────────────────────────────────────────────────────────
    _titulo(doc, "4. Integrações")
    _bullet(doc, "API do SEI", " — consulta e criação de processos, inclusão e download de documentos, controle de prazos e histórico de andamentos.")
    _bullet(doc, "Microsoft SharePoint", " — leitura das listas da fiscalização (designações, checklist e municípios).")
    _bullet(doc, "Microsoft Entra ID", " — autenticação corporativa dos usuários.")

    # ── Números ──────────────────────────────────────────────────────────────
    _titulo(doc, "5. Situação atual")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.style = "Light List Accent 1"
    cab0 = tabela.rows[0].cells
    for i, t in enumerate(("Indicador", "Situação")):
        r = cab0[i].paragraphs[0].add_run(t)
        r.bold = True
    for rotulo, valor in [
        ("Aplicação", "Backend e interface web concluídos e em uso para testes"),
        ("Testes automatizados", "166 testes (142 no backend e 24 na interface)"),
        ("Registros na consulta unificada", "18.868 relatórios e processos"),
        ("Automações agendadas", "Varredura do SEI, controle de prazos e sincronizações"),
        ("Unidades atendidas", "6 unidades do processamento"),
    ]:
        linha = tabela.add_row().cells
        linha[0].paragraphs[0].add_run(rotulo)
        linha[1].paragraphs[0].add_run(valor)

    # ── Ganhos ───────────────────────────────────────────────────────────────
    _titulo(doc, "6. Ganhos obtidos")
    _bullet(doc, "Menos trabalho manual", " — a instauração, que exigia várias etapas no SEI, passou a ser feita em uma única ação.")
    _bullet(doc, "Redução da dependência do SharePoint", " — os dados de referência foram migrados para banco próprio.")
    _bullet(doc, "Rastreabilidade", " — todas as ações e mudanças de fase ficam registradas para auditoria.")
    _bullet(doc, "Controle de prazos", " — acompanhamento automático, reduzindo risco de perda de prazo legal.")
    _bullet(doc, "Base para indicadores", " — dados estruturados permitem exportação para Power BI.")

    # ── Próximos passos ──────────────────────────────────────────────────────
    _titulo(doc, "7. Próximos passos")
    doc.add_paragraph(
        "A continuidade do desenvolvimento depende da manutenção da licença do Kiro. "
        "Estão previstos: conclusão dos textos padrão junto à área de negócio, publicação em "
        "servidor definitivo com banco corporativo, e ajustes decorrentes do uso pelos servidores."
    )

    # ── Encerramento ─────────────────────────────────────────────────────────
    doc.add_paragraph()
    fecho = doc.add_paragraph()
    r = fecho.add_run(
        "O aplicativo foi construído do zero com o apoio do Kiro, incluindo as integrações com o "
        "SEI e o SharePoint, a automação das fases processuais e a interface web. A manutenção da "
        "licença permite concluir as pendências e dar suporte à operação."
    )
    r.italic = True
    r.font.color.rgb = CINZA

    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)
    doc.save(caminho)
    return caminho


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera o relatório do projeto em Word.")
    parser.add_argument(
        "--saida",
        default=os.path.join(
            os.path.expanduser("~"), "Downloads",
            "Relatorio_Processamento_Sancionatorio_Kiro.docx",
        ),
        help="Caminho do .docx a gerar.",
    )
    args = parser.parse_args()
    caminho = gerar(args.saida)
    print(f"Documento gerado: {caminho}")


if __name__ == "__main__":
    main()
