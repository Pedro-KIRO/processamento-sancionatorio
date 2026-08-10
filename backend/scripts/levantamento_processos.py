"""Levantamento de processos nas 6 unidades SEI — exporta para Excel.

Busca todos os processos de 2026 e 2025 em cada unidade, paginando com cuidado
para não sobrecarregar a API (delay entre páginas).

Uso:
    cd backend/
    python scripts/levantamento_processos.py

Gera: levantamento_processos.xlsx na raiz do backend.
"""
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dotenv import load_dotenv
load_dotenv(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env"))

from app.integrations.sei.client import SeiClient, SeiSettings

UNIDADES = {
    "110053117": "ECV / EPIV",
    "110051045": "Peritos",
    "110051042": "Autoescola (CFC)",
    "110051044": "Desmontes",
    "110051043": "Instituições de ensino",
    "110053119": "Despachantes / Pátios",
}

DELAY_ENTRE_PAGINAS = 2  # segundos entre cada requisição (evitar sobrecarregar)
DELAY_ENTRE_UNIDADES = 5  # segundos entre unidades


def criar_sei_client() -> SeiClient:
    settings = SeiSettings(
        token_url=os.environ.get("SEI_TOKEN_URL", ""),
        client_id=os.environ.get("SEI_CLIENT_ID", os.environ.get("CLIENT_ID", "")),
        client_secret=os.environ.get("SEI_CLIENT_SECRET", os.environ.get("CLIENT_SECRET", "")),
        api_base=os.environ.get("SEI_API_BASE", ""),
        sigla_sistema=os.environ.get("SEI_SIGLA_SISTEMA", "CSDR_PROCESSAMENTO"),
        identificacao_servico=os.environ.get("SEI_IDENTIFICACAO_SERVICO", ""),
        trace_id=os.environ.get("SEI_TRACE_ID", "CSDR_PROCESSAMENTO_SEI"),
    )
    return SeiClient(settings, timeout=30, max_tentativas=1)


def buscar_processos_unidade(sei: SeiClient, id_unidade: str, nome_unidade: str) -> list[dict]:
    """Busca TODOS os processos de uma unidade, paginando."""
    todos = []
    start = 0

    print(f"\n{'='*60}")
    print(f"  Unidade: {nome_unidade} ({id_unidade})")
    print(f"{'='*60}")

    while True:
        print(f"  Página {start} (já coletados: {len(todos)})...", end=" ", flush=True)
        try:
            resp = sei.listar_processos(id_unidade, limit=500, start=start, tipo="T")
        except Exception as e:
            print(f"ERRO: {e}")
            break

        processos = resp.get("listaProcessos", [])
        if not processos:
            print("vazia (fim)")
            break

        todos.extend(processos)
        print(f"+{len(processos)} processos")

        if len(processos) < 500:
            break

        start += 1
        time.sleep(DELAY_ENTRE_PAGINAS)

    print(f"  Total na unidade: {len(todos)} processos")
    return todos


def filtrar_por_ano(processos: list[dict], ano: str) -> list[dict]:
    """Filtra processos pelo ano no número formatado (ex: /2026-)."""
    return [p for p in processos if f"/{ano}-" in (p.get("protocoloProcedimento") or "")]


def main():
    sei = criar_sei_client()
    print("=== Levantamento de Processos SEI (2026 + 2025) ===\n")

    resultados = []  # lista de dicts para o Excel

    for id_unidade, nome_unidade in UNIDADES.items():
        processos = buscar_processos_unidade(sei, id_unidade, nome_unidade)

        # Filtrar por ano
        de_2026 = filtrar_por_ano(processos, "2026")
        de_2025 = filtrar_por_ano(processos, "2025")

        print(f"  → 2026: {len(de_2026)} processos")
        print(f"  → 2025: {len(de_2025)} processos")
        print(f"  → Outros: {len(processos) - len(de_2026) - len(de_2025)} processos")

        for p in de_2026:
            resultados.append({
                "Unidade": nome_unidade,
                "ID Unidade": id_unidade,
                "Ano": 2026,
                "Número do Processo": p.get("protocoloProcedimento", ""),
                "idProcedimento": p.get("idProcedimento", ""),
                "Tipo": p.get("tipoProcesso", ""),
            })

        for p in de_2025:
            resultados.append({
                "Unidade": nome_unidade,
                "ID Unidade": id_unidade,
                "Ano": 2025,
                "Número do Processo": p.get("protocoloProcedimento", ""),
                "idProcedimento": p.get("idProcedimento", ""),
                "Tipo": p.get("tipoProcesso", ""),
            })

        time.sleep(DELAY_ENTRE_UNIDADES)

    # Exportar para Excel
    print(f"\n{'='*60}")
    print(f"  TOTAL GERAL: {len(resultados)} processos (2026 + 2025)")
    print(f"{'='*60}")

    try:
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()

        # Aba resumo
        ws_resumo = wb.active
        ws_resumo.title = "Resumo"
        ws_resumo.append(["Unidade", "Processos 2026", "Processos 2025", "Total"])
        for id_u, nome_u in UNIDADES.items():
            qtd_2026 = len([r for r in resultados if r["ID Unidade"] == id_u and r["Ano"] == 2026])
            qtd_2025 = len([r for r in resultados if r["ID Unidade"] == id_u and r["Ano"] == 2025])
            ws_resumo.append([nome_u, qtd_2026, qtd_2025, qtd_2026 + qtd_2025])

        total_2026 = len([r for r in resultados if r["Ano"] == 2026])
        total_2025 = len([r for r in resultados if r["Ano"] == 2025])
        ws_resumo.append(["TOTAL", total_2026, total_2025, total_2026 + total_2025])

        # Formatar cabeçalho resumo
        for cell in ws_resumo[1]:
            cell.font = Font(bold=True)

        # Aba com todos os processos
        ws_dados = wb.create_sheet("Processos")
        ws_dados.append(["Unidade", "ID Unidade", "Ano", "Número do Processo", "idProcedimento", "Tipo"])
        for r in resultados:
            ws_dados.append([r["Unidade"], r["ID Unidade"], r["Ano"], r["Número do Processo"], r["idProcedimento"], r["Tipo"]])

        for cell in ws_dados[1]:
            cell.font = Font(bold=True)

        # Ajustar largura das colunas
        for ws in [ws_resumo, ws_dados]:
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col) + 2
                ws.column_dimensions[col[0].column_letter].width = min(max_len, 40)

        arquivo = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend", "levantamento_processos.xlsx")
        # Salvar na pasta backend
        arquivo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "levantamento_processos.xlsx")
        wb.save(arquivo)
        print(f"\n  Arquivo salvo: {os.path.abspath(arquivo)}")

    except ImportError:
        # Sem openpyxl, salvar como CSV
        import csv
        arquivo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "levantamento_processos.csv")
        with open(arquivo, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=["Unidade", "ID Unidade", "Ano", "Número do Processo", "idProcedimento", "Tipo"])
            writer.writeheader()
            writer.writerows(resultados)
        print(f"\n  Arquivo salvo (CSV): {os.path.abspath(arquivo)}")
        print("  (Instale openpyxl para exportar como Excel: pip install openpyxl)")

    print("\nConcluído!")


if __name__ == "__main__":
    main()
