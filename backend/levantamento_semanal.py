"""Levantamento semanal: coleta processos de todas as unidades e acumula no Excel.

Cada execução adiciona um novo bloco com a data da semana. O Excel fica no
OneDrive para acesso compartilhado.

Uso manual:  python levantamento_semanal.py
Agendar:     Agendador de Tarefas do Windows → toda sexta-feira
"""
import os
import re
import sys
import time
from datetime import date

from dotenv import load_dotenv
load_dotenv()

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.integrations.sei.client import SeiClient, SeiSettings

# Caminho do Excel no OneDrive
CAMINHO_EXCEL = r"C:\Users\pedro.hsilva\OneDrive - PRODESP\levantamento_processos.xlsx"

UNIDADES = [
    ("110053117", "ECV-EPIV"),
    ("110051045", "Peritos"),
    ("110051042", "Autoescola"),
    ("110051044", "Desmontes"),
    ("110051043", "Inst. Ensino"),
    ("110053119", "Despachantes-Patios"),
]


def criar_sei_client() -> SeiClient:
    settings = SeiSettings(
        token_url=os.getenv("SEI_TOKEN_URL", ""),
        client_id=os.getenv("SEI_CLIENT_ID", ""),
        client_secret=os.getenv("SEI_CLIENT_SECRET", ""),
        api_base=os.getenv("SEI_API_BASE", ""),
        sigla_sistema=os.getenv("SEI_SIGLA_SISTEMA", "CSDR_PROCESSAMENTO"),
        identificacao_servico=os.getenv("SEI_IDENTIFICACAO_SERVICO", ""),
        trace_id=os.getenv("SEI_TRACE_ID", "CSDR_PROCESSAMENTO_SEI"),
    )
    return SeiClient(settings, timeout=60)


def listar_todos_processos(sei: SeiClient, id_unidade: str) -> list[dict]:
    """Lista TODOS os processos de uma unidade paginando."""
    todos = []
    start = 0
    while True:
        print(f"    Página {start} (coletados: {len(todos)})...")
        try:
            resp = sei.listar_processos(id_unidade, limit=500, start=start, tipo="T")
        except Exception as e:
            print(f"    ERRO: {e}")
            break
        processos = resp if isinstance(resp, list) else resp.get("listaProcessos", [])
        if not processos:
            break
        todos.extend(processos)
        if len(processos) < 500:
            break
        start += 1
        time.sleep(1)
    return todos


def extrair_ano(numero: str) -> str:
    match = re.search(r"/(\d{4})-", numero)
    return match.group(1) if match else "Outro"


def coletar_dados(sei) -> dict:
    """Coleta dados de todas as unidades."""
    resultado = {}
    for id_unidade, nome_unidade in UNIDADES:
        print(f"\n  === {nome_unidade} ({id_unidade}) ===")
        todos = listar_todos_processos(sei, id_unidade)
        print(f"  Total: {len(todos)}")

        por_ano: dict[str, int] = {}
        for p in todos:
            numero = p.get("protocoloProcedimento", "")
            ano = extrair_ano(numero)
            por_ano[ano] = por_ano.get(ano, 0) + 1

        resultado[nome_unidade] = {"total": len(todos), "por_ano": por_ano}
        time.sleep(2)

    return resultado


def atualizar_excel(resultado: dict, data_coleta: date):
    """Abre o Excel existente (ou cria novo) e adiciona o bloco da semana."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    # Abrir existente ou criar novo
    if os.path.exists(CAMINHO_EXCEL):
        wb = load_workbook(CAMINHO_EXCEL)
    else:
        wb = Workbook()

    # Usar aba "Histórico Semanal" (criar se não existir)
    nome_aba = "Historico Semanal"
    if nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
    else:
        if "Sheet" in wb.sheetnames:
            ws = wb["Sheet"]
            ws.title = nome_aba
        else:
            ws = wb.create_sheet(title=nome_aba)

    # Encontrar a próxima linha vazia
    proxima_linha = ws.max_row + 1 if ws.max_row and ws["A1"].value else 1

    # Se é a primeira vez, adicionar cabeçalho fixo na linha 1
    if proxima_linha == 1:
        pass  # Não usar cabeçalho fixo — cada bloco tem seu próprio

    # Adicionar separação (linha vazia se não é o primeiro bloco)
    if proxima_linha > 1:
        proxima_linha += 1  # linha vazia de separação

    # --- Cabeçalho do bloco ---
    data_fmt = data_coleta.strftime("%d/%m/%Y")
    ws.cell(row=proxima_linha, column=1, value=f"Levantamento — Semana de {data_fmt}")
    ws.cell(row=proxima_linha, column=1).font = Font(bold=True, size=12)
    ws.cell(row=proxima_linha, column=1).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.cell(row=proxima_linha, column=1).font = Font(bold=True, size=12, color="FFFFFF")
    # Mesclar o cabeçalho
    todos_anos = set()
    for dados in resultado.values():
        todos_anos.update(dados["por_ano"].keys())
    anos_ordenados = sorted(todos_anos, reverse=True)
    num_cols = 2 + len(anos_ordenados)
    ws.merge_cells(start_row=proxima_linha, start_column=1, end_row=proxima_linha, end_column=num_cols)
    proxima_linha += 1

    # --- Cabeçalho das colunas ---
    headers = ["Unidade", "Total"] + [f"{a}" for a in anos_ordenados]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=proxima_linha, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    proxima_linha += 1

    # --- Dados ---
    for nome, dados in resultado.items():
        ws.cell(row=proxima_linha, column=1, value=nome)
        ws.cell(row=proxima_linha, column=2, value=dados["total"])
        for i, ano in enumerate(anos_ordenados, start=3):
            ws.cell(row=proxima_linha, column=i, value=dados["por_ano"].get(ano, 0))
        proxima_linha += 1

    # --- Total geral ---
    ws.cell(row=proxima_linha, column=1, value="TOTAL")
    ws.cell(row=proxima_linha, column=1).font = Font(bold=True)
    ws.cell(row=proxima_linha, column=2, value=sum(d["total"] for d in resultado.values()))
    ws.cell(row=proxima_linha, column=2).font = Font(bold=True)
    for i, ano in enumerate(anos_ordenados, start=3):
        total_ano = sum(d["por_ano"].get(ano, 0) for d in resultado.values())
        ws.cell(row=proxima_linha, column=i, value=total_ano)
        ws.cell(row=proxima_linha, column=i).font = Font(bold=True)

    # Ajustar largura
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    for i in range(len(anos_ordenados)):
        ws.column_dimensions[chr(ord("C") + i)].width = 10

    # Salvar
    wb.save(CAMINHO_EXCEL)
    print(f"\nExcel atualizado em: {CAMINHO_EXCEL}")


def main():
    hoje = date.today()
    print(f"=== Levantamento Semanal — {hoje.strftime('%d/%m/%Y')} ===\n")

    sei = criar_sei_client()
    resultado = coletar_dados(sei)

    atualizar_excel(resultado, hoje)

    # Resumo no console
    print("\n=== RESUMO ===")
    for nome, dados in resultado.items():
        print(f"  {nome}: {dados['total']} processos")
    print(f"  TOTAL: {sum(d['total'] for d in resultado.values())}")
    print("\nConcluído!")


if __name__ == "__main__":
    main()
