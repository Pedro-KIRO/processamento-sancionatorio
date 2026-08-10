"""Levantamento: lista todos os processos nas 6 unidades do CPSAR.

Gera um Excel com tabelas formatadas (filtráveis), separado por unidade + resumo.
"""
import os
import re
import sys
import time

from dotenv import load_dotenv
load_dotenv()

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.integrations.sei.client import SeiClient, SeiSettings

UNIDADES = [
    ("110053117", "ECV-EPIV"),
    ("110051045", "Peritos"),
    ("110051042", "Autoescola"),
    ("110051044", "Desmontes"),
    ("110051043", "Inst. Ensino"),
    ("110053119", "Despachantes-Patios"),
]


def criar_sei_client() -> SeiClient:
    settings = SeiSettings(
        token_url=os.getenv("SEI_TOKEN_URL", ""),
        client_id=os.getenv("SEI_CLIENT_ID", ""),
        client_secret=os.getenv("SEI_CLIENT_SECRET", ""),
        api_base=os.getenv("SEI_API_BASE", ""),
        sigla_sistema=os.getenv("SEI_SIGLA_SISTEMA", "CSDR_PROCESSAMENTO"),
        identificacao_servico=os.getenv("SEI_IDENTIFICACAO_SERVICO", ""),
        trace_id=os.getenv("SEI_TRACE_ID", "CSDR_PROCESSAMENTO_SEI"),
    )
    return SeiClient(settings, timeout=60)


def listar_todos_processos(sei: SeiClient, id_unidade: str) -> list[dict]:
    """Lista TODOS os processos de uma unidade paginando com start."""
    todos = []
    start = 0

    while True:
        print(f"    Página {start} (já coletados: {len(todos)})...")
        try:
            resp = sei.listar_processos(id_unidade, limit=500, start=start, tipo="T")
        except Exception as e:
            print(f"    ERRO na página {start}: {e}")
            break

        processos = resp if isinstance(resp, list) else resp.get("listaProcessos", [])

        if not processos:
            break

        todos.extend(processos)

        if len(processos) < 500:
            break

        start += 1
        time.sleep(1)

    return todos


def extrair_ano(numero: str) -> str:
    """Extrai o ano do número SEI (ex: 140.00320970/2026-81 → 2026)."""
    match = re.search(r"/(\d{4})-", numero)
    return match.group(1) if match else "Outro"


def main():
    sei = criar_sei_client()
    resultado = {}

    for id_unidade, nome_unidade in UNIDADES:
        print(f"\n{'='*60}")
        print(f"=== Unidade: {nome_unidade} ({id_unidade}) ===")
        print(f"{'='*60}")

        todos = listar_todos_processos(sei, id_unidade)
        print(f"  Total de processos na unidade: {len(todos)}")

        # Separar por ano
        por_ano: dict[str, list[dict]] = {}
        for p in todos:
            numero = p.get("protocoloProcedimento", "")
            ano = extrair_ano(numero)
            por_ano.setdefault(ano, []).append(p)

        for ano in sorted(por_ano.keys(), reverse=True):
            print(f"    {ano}: {len(por_ano[ano])} processos")

        resultado[nome_unidade] = {
            "id_unidade": id_unidade,
            "total": len(todos),
            "processos": todos,
            "por_ano": por_ano,
        }

        time.sleep(2)

    print("\n\n=== Gerando Excel ===")
    gerar_excel(resultado)


def gerar_excel(resultado: dict):
    """Gera Excel com tabelas formatadas e filtráveis."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()

    # --- Aba de Resumo ---
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo.append(["Unidade", "Total na unidade"])

    # Coletar todos os anos que existem
    todos_anos = set()
    for dados in resultado.values():
        todos_anos.update(dados["por_ano"].keys())
    anos_ordenados = sorted(todos_anos, reverse=True)

    # Cabeçalho do resumo com anos dinâmicos
    ws_resumo.delete_rows(1)
    headers_resumo = ["Unidade", "Total"] + [f"Ano {a}" for a in anos_ordenados]
    ws_resumo.append(headers_resumo)

    for nome, dados in resultado.items():
        linha = [nome, dados["total"]]
        for ano in anos_ordenados:
            linha.append(len(dados["por_ano"].get(ano, [])))
        ws_resumo.append(linha)

    # Formatar como tabela
    ultima_col = chr(ord("A") + len(headers_resumo) - 1)
    ultima_linha = len(resultado) + 1
    ref_resumo = f"A1:{ultima_col}{ultima_linha}"
    tabela_resumo = Table(displayName="Resumo", ref=ref_resumo)
    tabela_resumo.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws_resumo.add_table(tabela_resumo)

    # Ajustar largura das colunas
    ws_resumo.column_dimensions["A"].width = 22
    ws_resumo.column_dimensions["B"].width = 12
    for i, _ in enumerate(anos_ordenados, start=3):
        ws_resumo.column_dimensions[chr(ord("A") + i - 1)].width = 12

    # --- Aba por unidade ---
    for nome, dados in resultado.items():
        nome_aba = nome[:31]
        ws = wb.create_sheet(title=nome_aba)
        ws.append(["Nº Processo SEI", "Ano", "Unidade"])

        for p in dados["processos"]:
            numero = p.get("protocoloProcedimento", "")
            ano = extrair_ano(numero)
            ws.append([numero, ano, nome])

        # Formatar como tabela
        num_linhas = len(dados["processos"]) + 1
        if num_linhas > 1:
            ref_tabela = f"A1:C{num_linhas}"
            nome_tabela = f"Proc_{nome.replace(' ', '').replace('-', '_').replace('.', '')}"[:30]
            tabela = Table(displayName=nome_tabela, ref=ref_tabela)
            tabela.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            ws.add_table(tabela)

        # Ajustar largura
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 22

    caminho = os.path.join(os.path.dirname(os.path.abspath(__file__)), "levantamento_processos.xlsx")
    wb.save(caminho)
    print(f"Excel salvo em: {caminho}")

    # Resumo no console
    print("\n=== RESUMO ===")
    print(f"{'Unidade':<22} {'Total':<8}", end="")
    for ano in anos_ordenados:
        print(f" {ano:<8}", end="")
    print()
    print("-" * (30 + 8 * len(anos_ordenados)))
    for nome, dados in resultado.items():
        print(f"{nome:<22} {dados['total']:<8}", end="")
        for ano in anos_ordenados:
            print(f" {len(dados['por_ano'].get(ano, [])):<8}", end="")
        print()


if __name__ == "__main__":
    main()
